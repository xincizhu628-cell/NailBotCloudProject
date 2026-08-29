import base64
import hashlib
import io
import json
import sqlite3
import sys
import time
from pathlib import Path

from admin_db import DB_PATH, ensure_admin_schema


PRODUCT_HEADERS = {
    "product_id": ["product_id", "productid", "sku", "id", "商品id", "商品编号", "产品编号"],
    "product_name": ["product_name", "productname", "name", "商品名称", "产品名称"],
    "product_type": ["product_type", "producttype", "type", "category", "商品类型", "产品类型"],
    "unit_price": ["unit_price", "unitprice", "price", "价格", "单价"],
    "product_info": ["product_info", "productinfo", "description", "detail", "商品介绍", "产品介绍", "详情"],
    "image_url": ["image_url", "imageurl", "image", "cover", "cover_url", "picture"],
    "image_base64": ["image_base64", "imagebase64", "base64"],
    "style_tags": ["style_tags", "styletags", "tags", "标签", "风格标签"],
    "nail_shape": ["nail_shape", "nailshape", "shape", "甲形", "甲型"],
    "bound_device_id": ["bound_device_id", "bound_device", "device", "device_id", "equip_id", "绑定设备", "绑定设备id", "设备编号"],
    "stock_quantity": ["stock_quantity", "stock", "quantity", "库存", "库存数量"],
    "stock_s": ["stock_s", "stocks", "s_stock", "stock_s_size", "s", "S库存", "S码库存"],
    "stock_m": ["stock_m", "stockm", "m_stock", "stock_m_size", "m", "M库存", "M码库存"],
    "stock_l": ["stock_l", "stockl", "l_stock", "stock_l_size", "l", "L库存", "L码库存"],
    "stock_xl": ["stock_xl", "stockxl", "xl_stock", "stock_xl_size", "xl", "XL库存", "XL码库存"],
    "on_delivery": ["on_delivery", "delivery", "shipping", "是否邮寄", "需要邮寄"],
    "pickup_method": ["pickup_method", "pickupmethod", "delivery_method", "fulfillment", "取货方式", "支持取货方式", "配送方式", "自取邮寄"],
    "is_featured": ["is_featured", "featured", "精品", "是否精品", "精选", "is_best"],
    "status": ["status", "状态"],
}

REWARD_HEADERS = {
    "reward_id": ["reward_id", "rewardid", "sku", "id", "奖励id", "奖品id", "商品编号"],
    "reward_name": ["reward_name", "rewardname", "name", "奖励名称", "奖品名称", "商品名称"],
    "reward_type": ["reward_type", "rewardtype", "type", "category", "奖励类型", "奖品类型"],
    "unit_point_cost": ["unit_point_cost", "pointcost", "points", "积分价格", "所需积分"],
    "reward_info": ["reward_info", "rewardinfo", "description", "detail", "奖励介绍", "奖品介绍", "详情"],
    "image_url": ["image_url", "imageurl", "image", "cover", "cover_url", "picture"],
    "image_base64": ["image_base64", "imagebase64", "base64"],
    "bound_device_id": ["bound_device_id", "bound_device", "device", "device_id", "equip_id", "绑定设备", "绑定设备id", "设备编号"],
    "on_delivery": ["on_delivery", "delivery", "shipping", "是否邮寄", "需要邮寄"],
    "pickup_method": ["pickup_method", "pickupmethod", "delivery_method", "fulfillment", "取货方式", "支持取货方式", "配送方式", "自取邮寄"],
    "is_featured": ["is_featured", "featured", "精品", "是否精品", "精选", "is_best"],
    "stock_quantity": ["stock_quantity", "stock", "quantity", "库存", "库存数量"],
    "stock_s": ["stock_s", "stocks", "s_stock", "stock_s_size", "s", "S库存", "S码库存"],
    "stock_m": ["stock_m", "stockm", "m_stock", "stock_m_size", "m", "M库存", "M码库存"],
    "stock_l": ["stock_l", "stockl", "l_stock", "stock_l_size", "l", "L库存", "L码库存"],
    "stock_xl": ["stock_xl", "stockxl", "xl_stock", "stock_xl_size", "xl", "XL库存", "XL码库存"],
    "status": ["status", "状态"],
}


def normalize_header(value):
    return str(value or "").strip().lower().replace(" ", "").replace("-", "_")


def pick(row, header_map, field):
    for name in header_map[field]:
        key = normalize_header(name)
        if key in row and row[key] not in (None, ""):
            return row[key]
    return None


def bool_int(value, default=0):
    if value in (None, ""):
        return default
    text = str(value).strip().lower()
    return 1 if text in {"1", "true", "yes", "y", "是", "需要", "邮寄"} else 0


def pickup_method(value, default="both"):
    if value in (None, ""):
        return default
    text = str(value).strip().lower()
    if text in {"pickup", "self", "self_pickup", "自取", "自提", "到店自取"}:
        return "pickup"
    if text in {"shipping", "mail", "post", "delivery", "邮寄", "配送", "快递"}:
        return "shipping"
    if text in {"both", "all", "pickup+shipping", "自取/邮寄", "自提/邮寄", "both取货"}:
        return "both"
    return default


def product_type(value, default="穿戴甲"):
    text = str(value or "").strip()
    key = text.lower().replace(" ", "_").replace("-", "_")
    if key in {"press_on_nail", "press_on", "presson", "wearable_nail"} or text == "穿戴甲":
        return "穿戴甲"
    if key in {"printed_nail", "printing_nail", "printable_nail"} or text == "打印甲":
        return "打印甲"
    if key in {"nail_accessory", "accessory", "accessories"} or text == "配件":
        return "配件"
    return default if not text else text


def number(value, default=0):
    try:
        if value in (None, ""):
            return default
        return float(value)
    except Exception:
        return default


def integer(value, default=0):
    return int(number(value, default))


def stock_sizes(row, header_map):
    stocks = {
        "stock_s": integer(pick(row, header_map, "stock_s"), 0),
        "stock_m": integer(pick(row, header_map, "stock_m"), 0),
        "stock_l": integer(pick(row, header_map, "stock_l"), 0),
        "stock_xl": integer(pick(row, header_map, "stock_xl"), 0),
    }
    stocks["stock_quantity"] = sum(stocks.values())
    return stocks


def data_url_to_bytes(data_url):
    if not data_url:
        return b""
    if "," in data_url:
        data_url = data_url.split(",", 1)[1]
    return base64.b64decode(data_url)


def image_payload(value):
    text = str(value or "").strip()
    if not text:
        return None
    mime_type = "image/png"
    encoded = text
    if text.startswith("data:") and "," in text:
        header, encoded = text.split(",", 1)
        if ";" in header:
            mime_type = header[5:].split(";", 1)[0] or mime_type
    raw = base64.b64decode(encoded)
    if not raw:
        return None
    return {
        "mime_type": mime_type,
        "data_url": f"data:{mime_type};base64,{encoded}",
        "sha256": hashlib.sha256(raw).hexdigest(),
    }


def normalize_image_assets(connection, rows, asset_type):
    for row in rows:
        payload = image_payload(row.get("image_base64"))
        if not payload:
            continue
        existing = connection.execute(
            "SELECT asset_id, url FROM assets WHERE sha256=? LIMIT 1",
            (payload["sha256"],),
        ).fetchone()
        if existing:
            asset_id, url = existing
        else:
            asset_id = f"asset_{asset_type}_{payload['sha256'][:18]}"
            url = ""
            connection.execute(
                """
                INSERT OR REPLACE INTO assets (asset_id, asset_type, mime_type, url, base64_data, sha256)
                VALUES (?, ?, ?, ?, ?, ?)
                """,
                (asset_id, asset_type, payload["mime_type"], url, payload["data_url"], payload["sha256"]),
            )
        row["image_url"] = url or f"/api/admin/product-image?itemKind=asset&id={asset_id}"
        row["image_base64"] = ""


def rows_from_xlsx(data_url):
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(data_url_to_bytes(data_url)), data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [normalize_header(cell) for cell in rows[0]]
    output = []
    for row in rows[1:]:
        item = {headers[index]: row[index] for index in range(min(len(headers), len(row))) if headers[index]}
        if any(value not in (None, "") for value in item.values()):
            output.append(item)
    return output


def product_from_row(row):
    now = int(time.time() * 1000)
    product_name = pick(row, PRODUCT_HEADERS, "product_name") or "Imported Product"
    product_id = pick(row, PRODUCT_HEADERS, "product_id") or f"P-{now}"
    item = {
        "product_id": str(product_id),
        "product_type": product_type(pick(row, PRODUCT_HEADERS, "product_type"), "穿戴甲"),
        "product_name": str(product_name),
        "unit_price": number(pick(row, PRODUCT_HEADERS, "unit_price"), 0),
        "product_info": str(pick(row, PRODUCT_HEADERS, "product_info") or ""),
        "image_url": str(pick(row, PRODUCT_HEADERS, "image_url") or ""),
        "image_base64": str(pick(row, PRODUCT_HEADERS, "image_base64") or ""),
        "style_tags": str(pick(row, PRODUCT_HEADERS, "style_tags") or ""),
        "nail_shape": str(pick(row, PRODUCT_HEADERS, "nail_shape") or ""),
        "bound_device_id": str(pick(row, PRODUCT_HEADERS, "bound_device_id") or ""),
        "on_delivery": bool_int(pick(row, PRODUCT_HEADERS, "on_delivery"), 1),
        "pickup_method": pickup_method(pick(row, PRODUCT_HEADERS, "pickup_method"), "both"),
        "is_featured": bool_int(pick(row, PRODUCT_HEADERS, "is_featured"), 0),
        "status": str(pick(row, PRODUCT_HEADERS, "status") or "active"),
    }
    item.update(stock_sizes(row, PRODUCT_HEADERS))
    return item


def reward_from_row(row):
    now = int(time.time() * 1000)
    reward_name = pick(row, REWARD_HEADERS, "reward_name") or "Imported Reward"
    reward_id = pick(row, REWARD_HEADERS, "reward_id") or f"R-{now}"
    item = {
        "reward_id": str(reward_id),
        "reward_name": str(reward_name),
        "reward_type": str(pick(row, REWARD_HEADERS, "reward_type") or "coupon"),
        "unit_point_cost": integer(pick(row, REWARD_HEADERS, "unit_point_cost"), 0),
        "reward_info": str(pick(row, REWARD_HEADERS, "reward_info") or ""),
        "image_url": str(pick(row, REWARD_HEADERS, "image_url") or ""),
        "image_base64": str(pick(row, REWARD_HEADERS, "image_base64") or ""),
        "bound_device_id": str(pick(row, REWARD_HEADERS, "bound_device_id") or ""),
        "on_delivery": bool_int(pick(row, REWARD_HEADERS, "on_delivery"), 0),
        "pickup_method": pickup_method(pick(row, REWARD_HEADERS, "pickup_method"), "pickup"),
        "is_featured": bool_int(pick(row, REWARD_HEADERS, "is_featured"), 0),
        "status": str(pick(row, REWARD_HEADERS, "status") or "active"),
    }
    item.update(stock_sizes(row, REWARD_HEADERS))
    return item


def insert_products(connection, rows):
    connection.executemany(
        """
        INSERT OR REPLACE INTO products (
          product_id, product_type, product_name, unit_price, product_info,
          image_url, image_base64, style_tags, nail_shape, bound_device_id, stock_quantity,
          stock_s, stock_m, stock_l, stock_xl, on_delivery, pickup_method, is_featured, status
        ) VALUES (
          :product_id, :product_type, :product_name, :unit_price, :product_info,
          :image_url, :image_base64, :style_tags, :nail_shape, :bound_device_id, :stock_quantity,
          :stock_s, :stock_m, :stock_l, :stock_xl, :on_delivery, :pickup_method, :is_featured, :status
        )
        """,
        rows,
    )


def insert_rewards(connection, rows):
    connection.executemany(
        """
        INSERT OR REPLACE INTO rewards (
          reward_id, reward_name, reward_type, unit_point_cost, reward_info,
          image_url, image_base64, bound_device_id, on_delivery, stock_quantity,
          stock_s, stock_m, stock_l, stock_xl, pickup_method, is_featured, status
        ) VALUES (
          :reward_id, :reward_name, :reward_type, :unit_point_cost, :reward_info,
          :image_url, :image_base64, :bound_device_id, :on_delivery, :stock_quantity,
          :stock_s, :stock_m, :stock_l, :stock_xl, :pickup_method, :is_featured, :status
        )
        """,
        rows,
    )


def validate_bound_devices(connection, rows):
    requested = sorted({str(row.get("bound_device_id") or "").strip() for row in rows if str(row.get("bound_device_id") or "").strip()})
    if not requested:
        return
    placeholders = ",".join("?" for _ in requested)
    existing = {
        str(row[0])
        for row in connection.execute(
            f"""
            SELECT equip_id
            FROM device_info
            WHERE equip_id IN ({placeholders})
              AND type IN ('主机', 'main_unit')
              AND COALESCE(status, 'active')='active'
            """,
            requested,
        ).fetchall()
    }
    if missing := [device for device in requested if device not in existing]:
        raise ValueError(f"设备不存在无法添加：{', '.join(missing)}")


def main():
    try:
        payload = json.loads(sys.stdin.read() or "{}")
        item_type = payload.get("itemType") or "product"
        source = payload.get("source") or "single"
        raw_rows = rows_from_xlsx(payload.get("xlsx")) if source == "xlsx" else [payload.get("item") or payload]

        if not raw_rows:
            raise ValueError("No product data found.")

        if item_type == "reward":
            rows = [reward_from_row({normalize_header(k): v for k, v in row.items()}) for row in raw_rows]
        else:
            rows = [product_from_row({normalize_header(k): v for k, v in row.items()}) for row in raw_rows]

        connection = sqlite3.connect(DB_PATH)
        try:
            ensure_admin_schema(connection)
            validate_bound_devices(connection, rows)
            normalize_image_assets(connection, rows, "reward-products" if item_type == "reward" else "products")
            insert_rewards(connection, rows) if item_type == "reward" else insert_products(connection, rows)
            connection.commit()
        finally:
            connection.close()

        print(json.dumps({
            "ok": True,
            "itemType": item_type,
            "imported": len(rows),
            "ids": [row.get("reward_id") or row.get("product_id") for row in rows],
        }, ensure_ascii=False))
    except Exception as exc:
        print(json.dumps({"ok": False, "error": str(exc)}, ensure_ascii=False))


if __name__ == "__main__":
    main()


